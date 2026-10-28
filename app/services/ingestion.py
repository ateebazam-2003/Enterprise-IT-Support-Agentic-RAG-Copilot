
from pathlib import Path
from typing import Iterable
from langchain_core.documents import Document
from docx import Document as DocxDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.document_loaders import PyPDFLoader,TextLoader


SUPPORTED=['.pdf','.txt','.md',".docs"]


def load_file(path:Path):
    suffix=path.suffix.lower()
    
    if suffix=='.pdf':
        return PyPDFLoader(str(path),encoding='utf-8').load()
    elif suffix in ['.txt','.md']:
        return TextLoader(str(path),encoding='utf-8').load
    elif suffix=='.docs':
        doc=DocxDocument(str(path))
        text='\n'.join(p.text for p in doc.paragraphs if p.text.strip())
        return [Document(page_content=text,metadata={"Source":str(path)})]
    else:
        raise ValueError(f"File not Supported:{suffix}")
    

def chunking(docs:Iterable[Document]):
    splitter=RecursiveCharacterTextSplitter(chunk_size=900,chunk_overlap=120,add_start_index=True)
    return  splitter.split_documents(list(docs))



